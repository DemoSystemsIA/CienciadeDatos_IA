# -*- coding: utf-8 -*-
"""
Validación contra un archivo de referencia ya aprobado por el comité
(p.ej. Resumen_NEW_VIP_Filtro_3.xlsx). Compara el campo FILTRO y, si existe,
la hoja Resumen_Persona. Sirve para probar que un cambio de reglas no rompió
nada: si el reporte sale con 0 discrepancias, el motor sigue siendo el mismo.
"""
from __future__ import annotations
import re

import pandas as pd
import openpyxl


def _norm(s):
    try:
        if s is None or pd.isna(s):
            return ''
    except (TypeError, ValueError):
        pass
    if isinstance(s, float) and s.is_integer():
        s = int(s)
    return re.sub(r'\s+', ' ', str(s)).strip().rstrip('…').strip()


def validar(df_p: pd.DataFrame, ruta_referencia: str) -> dict:
    wb = openpyxl.load_workbook(ruta_referencia, data_only=True)
    mine = {str(r['DNI']).strip(): r for _, r in df_p.iterrows()}
    out = {'archivo': ruta_referencia, 'hojas': wb.sheetnames,
           'filtro': None, 'resumen': None, 'detalles': []}

    # ---- FILTRO en la primera hoja ----
    ws = wb[wb.sheetnames[0]]
    hdr = [c.value for c in ws[1]]
    col_f = next((h for h in hdr if str(h).strip().upper() == 'FILTRO'), None)
    if col_f and 'DNI' in hdr:
        ok = tot = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            d = dict(zip(hdr, row))
            dni = str(d.get('DNI') or '').strip()
            if not dni:
                continue
            tot += 1
            esp = _norm(d.get(col_f))
            got = _norm(mine[dni]['FILTRO']) if dni in mine else '(DNI no encontrado)'
            if esp == got:
                ok += 1
            else:
                out['detalles'].append({'HOJA': ws.title, 'DNI': dni, 'CAMPO': 'FILTRO',
                                        'ESPERADO': esp, 'OBTENIDO': got})
        out['filtro'] = {'ok': ok, 'total': tot}

    # ---- Resumen_Persona ----
    if 'Resumen_Persona' in wb.sheetnames:
        ws = wb['Resumen_Persona']
        hdr = [c.value for c in ws[1]]
        mapa = {'NIVEL DE RIESGO': 'NIVEL_RIESGO', 'GRAVEDAD MAXIMA': 'GRAVEDAD_MAX',
                'AÑO MAS RECIENTE': 'ANIO_MAX', 'VEREDICTO': 'VEREDICTO',
                'INDICE RIESGO (0-100)': 'INDICE', 'IDX GRAVEDAD': 'IDX_GRAVEDAD',
                'IDX VIGENCIA': 'IDX_VIGENCIA', 'IDX REINCIDENCIA': 'IDX_REINCIDENCIA',
                'IDX CASO ACTIVO': 'IDX_CASO_ACTIVO', 'IDX DISPERSION': 'IDX_DISPERSION',
                'CATEGORIAS': 'CATEGORIAS',
                'INCIDENCIAS POLICIALES (DETALLE 1)': 'INCID_POLICIALES'}
        ok = tot = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            d = dict(zip(hdr, row))
            dni = str(d.get('DNI') or '').strip()
            if not dni or dni not in mine:
                continue
            m = mine[dni]
            for col_ref, col_mio in mapa.items():
                if col_ref not in hdr or col_mio not in df_p.columns:
                    continue
                tot += 1
                a = _norm(d.get(col_ref))
                b = _norm(m[col_mio])
                if a in ('', 'None') and b in ('', 'None', 'nan'):
                    ok += 1
                elif a == b:
                    ok += 1
                else:
                    out['detalles'].append({'HOJA': 'Resumen_Persona', 'DNI': dni,
                                            'CAMPO': col_ref, 'ESPERADO': a, 'OBTENIDO': b})
        out['resumen'] = {'ok': ok, 'total': tot}
    return out
