# -*- coding: utf-8 -*-
"""
Recorre la carpeta raíz, consolida los DNI de todos los Resumen_NEW_VIP_*.xlsx,
los cruza con los PDF de cada subcarpeta Adjuntos y arma los DataFrames.

Cachea el parseo de cada PDF en disco (.cache_filtro) usando ruta+mtime+tamaño
como llave, así la segunda corrida es instantánea aunque reinicies Streamlit.
"""
from __future__ import annotations
import os
import re
import pickle
import hashlib
from collections import OrderedDict, Counter

import pandas as pd
import openpyxl

from . import config as C
from . import maestro as M
from .pdf_reader import parse_pdf
from .classify import (clasificar, gravedad, caso_activo, anio, labor_personas,
                       indice_riesgo, nivel_riesgo, veredicto)

CACHE_DIR = ".cache_filtro"
RE_DNI = re.compile(r'(\d{6,})')


# ------------------------------------------------------------ escaneo
def escanear(raiz: str) -> dict:
    """Devuelve {'excels': [...], 'pdfs': [...], 'carpetas': [...]} con rutas absolutas."""
    raiz = os.path.abspath(raiz)
    excels, pdfs, carpetas = [], [], []
    if not os.path.isdir(raiz):
        return {'excels': [], 'pdfs': [], 'carpetas': [], 'raiz': raiz, 'existe': False}
    for nombre in sorted(os.listdir(raiz)):
        sub = os.path.join(raiz, nombre)
        if not os.path.isdir(sub) or nombre.startswith('.'):
            continue
        import glob
        ex = sorted(glob.glob(os.path.join(sub, C.PATRON_EXCEL)))
        pd_ = sorted(glob.glob(os.path.join(sub, C.SUBCARPETA_PDF, C.PATRON_PDF)))
        if not ex and not pd_:
            continue
        carpetas.append(nombre)
        excels += [(nombre, p) for p in ex if not os.path.basename(p).startswith('~$')]
        pdfs += [(nombre, p) for p in pd_]

    # Reportes comunes: una carpeta «Adjuntos» colgando de la raíz vale para todos
    # los DNI (el cruce PDF -> persona es por DNI, no por carpeta). Es lo que usa
    # la versión web cuando se suben los archivos sueltos, sin estructura.
    import glob as _glob
    comunes = sorted(_glob.glob(os.path.join(raiz, C.SUBCARPETA_PDF, C.PATRON_PDF)))
    pdfs += [("(común)", p) for p in comunes]

    return {'excels': excels, 'pdfs': pdfs, 'carpetas': carpetas, 'raiz': raiz,
            'maestro': M.localizar(raiz), 'existe': True}


def firma(scan: dict) -> str:
    """Huella de la carpeta: cambia si se agrega, quita o modifica cualquier archivo."""
    h = hashlib.sha1()
    sueltos = [(None, scan['maestro'])] if scan.get('maestro') else []
    for _, p in scan['excels'] + scan['pdfs'] + sueltos:
        try:
            st = os.stat(p)
            h.update(f"{p}|{int(st.st_mtime)}|{st.st_size}".encode('utf-8', 'ignore'))
        except OSError:
            h.update(p.encode('utf-8', 'ignore'))
    return h.hexdigest()


# ----------------------------------------------------- caché de PDFs
def _cache_path(raiz: str) -> str:
    d = os.path.join(raiz, CACHE_DIR)
    os.makedirs(d, exist_ok=True)
    return os.path.join(d, "pdfs.pkl")


def _cargar_cache(raiz):
    try:
        with open(_cache_path(raiz), 'rb') as f:
            return pickle.load(f)
    except Exception:
        return {}


def _guardar_cache(raiz, data):
    try:
        with open(_cache_path(raiz), 'wb') as f:
            pickle.dump(data, f)
    except Exception:
        pass


def parsear_pdfs(scan: dict, progreso=None) -> dict:
    """{dni: resultado_parse} tomando, si un DNI se repite, el reporte más completo."""
    raiz = scan['raiz']
    cache = _cargar_cache(raiz)
    nuevo = dict(cache)
    por_dni, total = {}, len(scan['pdfs'])

    for i, (_carpeta, path) in enumerate(scan['pdfs'], 1):
        m = RE_DNI.search(os.path.basename(path))
        if not m:
            continue
        dni = m.group(1)
        try:
            st = os.stat(path)
            key = f"{path}|{int(st.st_mtime)}|{st.st_size}"
        except OSError:
            key = path
        res = cache.get(key)
        if res is None:
            res = parse_pdf(path)
            nuevo[key] = res
        res = dict(res)
        res['path'] = path
        prev = por_dni.get(dni)
        if prev is None or len(res['recs']) > len(prev['recs']):
            res['n_pdf'] = (prev or {}).get('n_pdf', 0) + 1
            por_dni[dni] = res
        else:
            prev['n_pdf'] = prev.get('n_pdf', 1) + 1
        if progreso:
            progreso(i, total, os.path.basename(path))

    _guardar_cache(raiz, nuevo)
    return por_dni


# ----------------------------------------------------------- consolidado
def _limpiar_jurisdiccion(txt: str) -> str:
    txt = re.sub(r'^(Sede\s*D\.?\s*J\.?|Sede|DF|Oficina)\b', '', txt or '', flags=re.I)
    return re.sub(r'^[\s\.,;:\-]+', '', txt).strip().title()


def construir(raiz: str, progreso=None):
    """
    Devuelve (df_personas, df_delitos, meta).
    df_personas: una fila por DNI único. df_delitos: una fila por delito en rojo.
    """
    scan = escanear(raiz)
    if not scan['existe']:
        raise FileNotFoundError(f"No existe la carpeta: {raiz}")
    if not scan['excels']:
        raise FileNotFoundError(
            f"No se encontró ningún '{C.PATRON_EXCEL}' dentro de las subcarpetas de {raiz}")

    # ---- 1. DNI desde los Excel ----
    personas, filas_excel = OrderedDict(), 0
    for carpeta, path in scan['excels']:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb[C.HOJA_EXCEL] if C.HOJA_EXCEL in wb.sheetnames else wb[wb.sheetnames[0]]
        filas = ws.iter_rows(values_only=True)
        hdr = [c for c in next(filas)]
        for r in filas:
            d = dict(zip(hdr, r))
            dni = str(d.get(C.COL_DNI) or '').strip()
            if not dni or dni.lower() == 'none':
                continue
            filas_excel += 1
            base = personas.setdefault(dni, {**d, '_CARPETAS': [], '_ARCHIVOS': []})
            for k, v in d.items():                       # completa huecos entre archivos
                if base.get(k) in (None, '') and v not in (None, ''):
                    base[k] = v
            if carpeta not in base['_CARPETAS']:
                base['_CARPETAS'].append(carpeta)
            nb = os.path.basename(path)
            if nb not in base['_ARCHIVOS']:
                base['_ARCHIVOS'].append(nb)
        wb.close()

    # ---- 2. PDFs ----
    parsed = parsear_pdfs(scan, progreso=progreso)

    # ---- 3. Cruce ----
    filas_delitos, filas_personas = [], []
    for dni, d in personas.items():
        nombre = str(d.get(C.COL_NOMBRE) or '').strip()
        labor, tipo = labor_personas(d['_CARPETAS'])
        pr = parsed.get(dni)
        tiene_pdf = pr is not None
        delitos = []

        for rec in (pr['recs'] if pr else []):
            texto = (rec.get('Delito') or '').strip()
            if not texto:
                continue
            niv, cat, regla = clasificar(texto)
            a = anio(rec)
            estado = rec.get('Estado') or (
                'Registro policial (sin estado)'
                if 'DET' in (rec.get('Especialidad') or '') else '')
            fuente = ('Carpeta fiscal (MP)' if 'Estado' in rec else
                      'Detención policial (RENADESPPLE)'
                      if ('Motivo Detención' in rec or 'Fecha Detención' in rec)
                      else 'Registro policial')
            jur = _limpiar_jurisdiccion(rec.get('Jurisdicción') or rec.get('Distrito Fiscal')
                                        or rec.get('Distrito') or '')
            g = gravedad(cat)
            delitos.append(dict(
                DNI=dni, NOMBRE=nombre, LABOR=labor, TIPO_PERSONAL=tipo,
                DELITO=texto, NIVEL=niv, CONCEPTO=C.NIVELES[niv][0],
                ACCION=C.NIVELES[niv][1], CATEGORIA=cat, GRAVEDAD=g,
                PESO_GRAVEDAD=C.PESO_GRAV[g], ANIO=a,
                VIGENTE=bool(a and a >= C.ANIO_VIGENCIA),
                ACTIVO=caso_activo(estado), JURISDICCION=jur, FUENTE=fuente,
                ESTADO=(estado or 'Sin estado'), PARTE=rec.get('Parte', ''),
                ESPECIALIDAD=rec.get('Especialidad', ''),
                ENTIDAD=(rec.get('Dependencia') or rec.get('Fiscalía')
                         or rec.get('Dependencia Policial')
                         or rec.get('Entidad Informante') or ''),
                CASO=rec.get('Caso', ''), REGLA=regla))

        # OJO: los delitos se conservan en el ORDEN DEL REPORTE.
        # El campo FILTRO debe respetarlo (así lo valida el archivo de referencia);
        # el orden por criticidad se aplica solo al mostrar, nunca aquí.
        for i, x in enumerate(delitos, 1):
            x['ORDEN_REPORTE'] = i
        filas_delitos += delitos
        n = len(delitos)

        if not tiene_pdf:
            filtro = 'SIN PDF'
        elif n == 0:
            filtro = 'SIN DELITOS EN ROJO'
        else:
            filtro = ' || '.join(x['DELITO'] for x in delitos)

        ind, idx, grav, amax = indice_riesgo(delitos)
        nr = nivel_riesgo(grav, amax, tiene_pdf, n)
        ver = veredicto(nr, tiene_pdf, n)
        nivmin = min((x['NIVEL'] for x in delitos), default=None)
        juris = sorted({x['JURISDICCION'] for x in delitos if x['JURISDICCION']})

        fila = {k: v for k, v in d.items() if not k.startswith('_') and k is not None}
        fila.update(dict(
            DNI=dni, **{C.COL_NOMBRE: nombre},
            LABOR=labor, TIPO_PERSONAL=tipo,
            CARPETAS=", ".join(d['_CARPETAS']), CARPETAS_L=d['_CARPETAS'],
            ARCHIVOS=", ".join(d['_ARCHIVOS']),
            TIENE_PDF=tiene_pdf, PDF_ADJUNTO="SI" if tiene_pdf else "NO",
            FILTRO=filtro, N_DELITOS=n,
            NIVEL_MATRIZ=(f"N{nivmin}" if nivmin else ""), NIVEL_NUM=nivmin,
            CONCEPTO_MATRIZ=(C.NIVELES[nivmin][0] if nivmin else ""),
            ACCION_MATRIZ=(C.NIVELES[nivmin][1] if nivmin
                           else ("SOLICITAR REPORTE" if not tiene_pdf else "NINGUNA")),
            NIVEL_RIESGO=nr, GRAVEDAD_MAX=grav,
            CATEGORIAS=", ".join(sorted({x['CATEGORIA'] for x in delitos})),
            CATEGORIAS_L=sorted({x['CATEGORIA'] for x in delitos}),
            ANIO_MAX=amax, VIGENTE="SI" if (amax and amax >= C.ANIO_VIGENCIA) else "NO",
            CASOS_ACTIVOS=sum(1 for x in delitos if x['ACTIVO']),
            DELITOS_VIGENTES=sum(1 for x in delitos if x['VIGENTE']),
            JURISDICCIONES=" | ".join(juris), N_JURISDICCIONES=len(juris),
            JURIS_L=juris,
            INCID_POLICIALES=(pr['incid'] if pr else 0),
            PROVINCIA=(pr['prov'] if pr else ""),
            VEREDICTO=ver, INDICE=ind,
            IDX_GRAVEDAD=idx[0], IDX_VIGENCIA=idx[1], IDX_REINCIDENCIA=idx[2],
            IDX_CASO_ACTIVO=idx[3], IDX_DISPERSION=idx[4],
        ))
        filas_personas.append(fila)

    df_p = pd.DataFrame(filas_personas)
    df_d = pd.DataFrame(filas_delitos)
    if df_d.empty:
        df_d = pd.DataFrame(columns=['DNI', 'NOMBRE', 'LABOR', 'TIPO_PERSONAL', 'DELITO',
                                     'NIVEL', 'CONCEPTO', 'ACCION', 'CATEGORIA', 'GRAVEDAD',
                                     'PESO_GRAVEDAD', 'ANIO', 'VIGENTE', 'ACTIVO',
                                     'JURISDICCION', 'FUENTE', 'ESTADO', 'PARTE',
                                     'ESPECIALIDAD', 'ENTIDAD', 'CASO', 'REGLA',
                                     'ORDEN_REPORTE'])

    # enteros anulables: evita que pandas convierta los años a float por los nulos
    for col in ('ANIO_MAX', 'NIVEL_NUM'):
        if col in df_p.columns:
            df_p[col] = pd.array(df_p[col], dtype='Int64')
    if 'ANIO' in df_d.columns and len(df_d):
        df_d['ANIO'] = pd.array(df_d['ANIO'], dtype='Int64')

    # ---- 4. Cruce con el maestro de funcionarios de Prize ----
    ruta_m = scan.get('maestro')
    maestro_df, err_maestro = None, ""
    if ruta_m:
        try:
            maestro_df = M.cargar(ruta_m)
        except Exception as e:                      # archivo raro: seguimos sin él
            err_maestro = f"{os.path.basename(ruta_m)}: {e}"
    df_p, info_maestro = M.aplicar(df_p, maestro_df)
    info_maestro['archivo'] = os.path.basename(ruta_m) if ruta_m else ""
    info_maestro['error'] = err_maestro

    orden = {v: i for i, v in enumerate(C.VEREDICTOS)}
    df_p = df_p.sort_values(
        by=['VEREDICTO', 'INDICE', C.COL_NOMBRE],
        key=lambda s: s.map(orden) if s.name == 'VEREDICTO' else (-s if s.name == 'INDICE' else s)
    ).reset_index(drop=True)

    dnis_pdf = set(parsed)
    meta = dict(
        raiz=scan['raiz'], carpetas=scan['carpetas'],
        n_excels=len(scan['excels']), n_pdfs=len(scan['pdfs']),
        filas_excel=filas_excel, dni_unicos=len(personas),
        dni_repetidos=filas_excel - len(personas),
        con_pdf=int(df_p['TIENE_PDF'].sum()), sin_pdf=int((~df_p['TIENE_PDF']).sum()),
        pdfs_huerfanos=sorted(dnis_pdf - set(personas)),
        maestro=info_maestro,
        firma=firma(scan),
    )
    return df_p, df_d, meta
